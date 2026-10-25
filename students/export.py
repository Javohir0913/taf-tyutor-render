"""Talabalar ro'yxatini Excel (.xlsx) fayliga eksport qilish."""
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    # (sarlavha, ustun kengligi, qiymat olish funksiyasi)
    ("№", 6, None),
    ("Familiya", 18, lambda s: s.last_name),
    ("Ism", 16, lambda s: s.first_name),
    ("Otasining ismi", 18, lambda s: s.middle_name),
    ("Kurs", 8, lambda s: f"{s.course}-kurs"),
    ("Guruh", 12, lambda s: s.group.name),
    ("Yashash joyi", 16, lambda s: s.get_residence_type_display()),
    ("Hozirgi manzil", 42, lambda s: s.current_address),
]


def students_to_xlsx(queryset):
    wb = Workbook()
    ws = wb.active
    ws.title = "Talabalar"

    header_fill = PatternFill("solid", fgColor="0F6B5A")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFC9C5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    for col, (title, width, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.fill, c.font, c.border = header_fill, header_font, border
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    qs = queryset.select_related("group").order_by("course", "group__name", "last_name", "first_name")
    for i, s in enumerate(qs, start=1):
        ws.cell(row=i + 1, column=1, value=i)
        for col, (_, _, getter) in enumerate(COLUMNS[1:], start=2):
            ws.cell(row=i + 1, column=col, value=getter(s))
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=i + 1, column=col)
            cell.border, cell.alignment = border, wrap

    ws.freeze_panes = "A2"  # sarlavha qatori qotib turadi
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(queryset):
    filename = f"talabalar_{timezone.localtime():%Y-%m-%d_%H-%M}.xlsx"
    response = HttpResponse(
        students_to_xlsx(queryset),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
